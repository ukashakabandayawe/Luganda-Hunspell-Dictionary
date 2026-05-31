from __future__ import annotations

import argparse
import re
from collections import defaultdict
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook


HEADER_RE = re.compile(r"^#\s*-+\s*(.*?)\s*-+\s*#\s*$")
FOOTER_RE = re.compile(r"^#\s*-+\s*#\s*$")
DEFAULT_WORKBOOK = Path("resources/learn-Luganda Dictionary.xlsx")
DEFAULT_DICTIONARY = Path("Luganda.dic")
DEFAULT_HEADER_TOTAL_WIDTH = 108
DEFAULT_FOOTER_TOTAL_WIDTH = 109


@dataclass(frozen=True)
class Candidate:
    stem: str
    meaning: str
    pos: str
    sheet: str
    row_no: int


def normalize(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def is_target_pos(pos_text: str) -> bool:
    lowered = pos_text.lower()
    return "verb" in lowered or "adj" in lowered


def has_whitespace(text: str) -> bool:
    return any(ch.isspace() for ch in text)


def parse_entry_line(line: str) -> tuple[str, str, str] | None:
    if line is None:
        return None

    trimmed = line.strip()
    if not trimmed or trimmed.startswith("#"):
        return None

    first_whitespace = -1
    for index, char in enumerate(line):
        if char.isspace():
            first_whitespace = index
            break

    token = line[:first_whitespace] if first_whitespace >= 0 else line
    trailing = line[first_whitespace:] if first_whitespace >= 0 else ""
    slash = token.find("/")
    stem = token[:slash] if slash >= 0 else token
    flags_raw = token[slash + 1 :] if slash >= 0 else ""
    stem = stem.strip()

    if not stem:
        return None

    return stem, flags_raw, trailing


def format_header(stem: str, meaning: str, total_width: int = DEFAULT_HEADER_TOTAL_WIDTH) -> str:
    title = f"oku{stem}({meaning})"
    prefix = "# "
    suffix = "#"
    dash_space_budget = max(total_width - len(prefix) - len(suffix) - len(title) - 2, 0)
    left_dashes = dash_space_budget // 2
    right_dashes = dash_space_budget - left_dashes
    return f"{prefix}{'-' * left_dashes} {title} {'-' * right_dashes}{suffix}"


def format_footer(total_width: int = DEFAULT_FOOTER_TOTAL_WIDTH) -> str:
    if total_width < 5:
        total_width = 5
    return f"# {'-' * (total_width - 4)} #"


def collect_candidates(workbook_path: Path) -> tuple[dict[str, Candidate], list[str]]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    candidates: dict[str, Candidate] = {}
    conflicts: list[str] = []

    for worksheet in workbook.worksheets:
        for row_no in range(1, worksheet.max_row + 1):
            stem = normalize(worksheet.cell(row_no, 1).value)
            pos = normalize(worksheet.cell(row_no, 2).value)
            meaning = normalize(worksheet.cell(row_no, 4).value)

            if not stem or not pos or not meaning:
                continue
            if has_whitespace(stem):
                continue
            if not is_target_pos(pos):
                continue

            existing = candidates.get(stem)
            if existing is None:
                candidates[stem] = Candidate(stem=stem, meaning=meaning, pos=pos, sheet=worksheet.title, row_no=row_no)
                continue

            # Combine differing meanings for the same stem (preserve order, deduplicate)
            if existing.meaning != meaning:
                existing_parts = [p.strip() for p in existing.meaning.split(" / ") if p.strip()]
                if meaning not in existing_parts:
                    combined_parts = existing_parts + [meaning]
                    combined_meaning = " / ".join(combined_parts)
                    # replace with a new Candidate carrying the combined meaning and merged pos
                    combined_pos = ",".join(p for p in [existing.pos, pos] if p)
                    candidates[stem] = Candidate(
                        stem=stem,
                        meaning=combined_meaning,
                        pos=combined_pos,
                        sheet=f"{existing.sheet},{worksheet.title}",
                        row_no=existing.row_no,
                    )
                    conflicts.append(
                        f"{stem}: combined meanings -> {existing.meaning!r} + {meaning!r}"
                    )

    return candidates, conflicts


def scan_dictionary(dictionary_path: Path) -> tuple[list[str], set[str], dict[str, int]]:
    lines = dictionary_path.read_text(encoding="utf-8").splitlines(keepends=True)
    inside_group = False
    grouped_stems: set[str] = set()
    free_stem_lines: dict[str, int] = {}

    for index, line in enumerate(lines):
        stripped = line.strip()
        if not stripped:
            continue

        if stripped.startswith("#"):
            if FOOTER_RE.match(stripped):
                inside_group = False
            elif HEADER_RE.match(stripped):
                inside_group = True
            continue

        entry = parse_entry_line(line)
        if entry is None:
            continue

        stem, _, _ = entry
        if inside_group:
            grouped_stems.add(stem)
            continue

        if stem not in free_stem_lines:
            free_stem_lines[stem] = index
        else:
            free_stem_lines[stem] = -1

    return lines, grouped_stems, free_stem_lines


def build_annotated_lines(
    lines: list[str],
    candidates: dict[str, Candidate],
    grouped_stems: set[str],
    free_stem_lines: dict[str, int],
) -> tuple[list[str], list[str], list[str]]:
    insert_before: dict[int, list[str]] = defaultdict(list)
    insert_after: dict[int, list[str]] = defaultdict(list)
    applied: list[str] = []
    skipped: list[str] = []

    for stem, candidate in candidates.items():
        if stem in grouped_stems:
            skipped.append(f"{stem}: already covered by an existing manual group")
            continue

        line_index = free_stem_lines.get(stem)
        if line_index is None:
            skipped.append(f"{stem}: not found in Luganda.dic")
            continue
        if line_index < 0:
            skipped.append(f"{stem}: appears multiple times outside manual groups")
            continue

        insert_before[line_index].append(format_header(candidate.stem, candidate.meaning))
        insert_after[line_index].append(format_footer())
        applied.append(stem)

    output_lines: list[str] = []
    for index, line in enumerate(lines):
        for header_line in insert_before.get(index, []):
            output_lines.append(header_line + "\n")
        output_lines.append(line)
        for footer_line in insert_after.get(index, []):
            output_lines.append(footer_line + "\n")

    return output_lines, applied, skipped


def main() -> int:
    parser = argparse.ArgumentParser(description="Annotate Luganda.dic stem groups from the workbook.")
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--dictionary", type=Path, default=DEFAULT_DICTIONARY)
    parser.add_argument("--verbs-only", action="store_true", help="Only process candidates whose POS contains 'verb'.")
    parser.add_argument("--output", type=Path, help="Write annotated output to this path.")
    parser.add_argument("--preview", nargs="*", default=[], help="Show generated wrapper lines for these stems.")
    args = parser.parse_args()

    candidates, conflicts = collect_candidates(args.workbook)
    if args.verbs_only:
        candidates = {s: c for s, c in candidates.items() if "verb" in (c.pos or "").lower()}
    lines, grouped_stems, free_stem_lines = scan_dictionary(args.dictionary)
    output_lines, applied, skipped = build_annotated_lines(lines, candidates, grouped_stems, free_stem_lines)

    print(f"Workbook candidates: {len(candidates)}")
    print(f"Manual-group stems skipped: {len(grouped_stems)}")
    print(f"Applied wrappers: {len(applied)}")
    print(f"Skipped candidates: {len(skipped)}")

    if conflicts:
        print("Conflicts:")
        for conflict in conflicts[:20]:
            print(f"  {conflict}")

    if args.preview:
        print("Preview:")
        for stem in args.preview:
            candidate = candidates.get(stem)
            if candidate is None:
                print(f"  {stem}: not found in workbook")
                continue
            print(format_header(candidate.stem, candidate.meaning))
            print(stem)
            print(format_footer())

    if args.output:
        outpath: Path = args.output
        # If writing over the original dictionary, make a backup first
        try:
            if outpath.resolve() == args.dictionary.resolve():
                original_text = args.dictionary.read_text(encoding="utf-8")
                backup_path = args.dictionary.with_name(args.dictionary.name + ".bak")
                backup_path.write_text(original_text, encoding="utf-8")
                print(f"Created backup of original dictionary at {backup_path}")
        except Exception:
            # fail-safe: do not block writing if backup creation fails
            print("Warning: failed to create backup; proceeding to write output")

        outpath.write_text("".join(output_lines), encoding="utf-8")
        print(f"Wrote annotated dictionary to {outpath}")

    return 0


if __name__ == "__main__":
    raise SystemExit(main())